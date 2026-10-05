from __future__ import annotations

from pathlib import Path
import argparse
from docx import Document
from docx.shared import Inches


def find_caption(doc: Document, startswith: str):
    for p in doc.paragraphs:
        if p.text.strip().startswith(startswith):
            return p
    raise ValueError(f"Caption not found: {startswith}")


def preceding_drawing_paragraph(caption):
    node = caption._p.getprevious()
    while node is not None:
        if node.xpath(".//w:drawing"):
            return node
        node = node.getprevious()
    raise ValueError(f"No preceding drawing found for {caption.text[:40]}")


def replace_drawing(doc: Document, caption_prefix: str, image_path: Path, width_inches: float = 6.45):
    caption = find_caption(doc, caption_prefix)
    drawing_p = preceding_drawing_paragraph(caption)

    # Remove existing runs/drawings but preserve paragraph properties.
    for child in list(drawing_p):
        if child.tag.endswith("}pPr"):
            continue
        drawing_p.remove(child)

    # python-docx paragraph wrapper for the same XML node.
    from docx.text.paragraph import Paragraph
    p = Paragraph(drawing_p, caption._parent)
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))


def main():
    parser = argparse.ArgumentParser(
        description="Replace corrected Figures 3 and 5 in the manuscript after regeneration from Bronze data."
    )
    parser.add_argument("input_docx", type=Path)
    parser.add_argument("output_docx", type=Path)
    parser.add_argument(
        "--figures-dir",
        type=Path,
        default=Path("outputs/figures_q1_final_complete"),
    )
    args = parser.parse_args()

    figures = args.figures_dir.resolve()
    fig3 = figures / "Fig03_bronze_overview_CORRECTED.png"
    fig5 = figures / "Fig05_high_frequency_returns_CORRECTED.png"
    for p in (fig3, fig5):
        if not p.exists():
            raise FileNotFoundError(p)

    doc = Document(args.input_docx)
    replace_drawing(doc, "Fig. 3 Bronze one-minute dataset overview.", fig3)
    replace_drawing(doc, "Fig. 5 High-frequency return behaviour", fig5)

    c3 = find_caption(doc, "Fig. 3 Bronze one-minute dataset overview.")
    c3.text = (
        "Fig. 3 Bronze one-minute dataset overview. (a) BTCUSDT closing-price history; "
        "(b) closing-price distribution; (c) one-minute trading volume; and (d) number "
        "of trades per minute. Panels (c) and (d) use linear x-axes limited to the "
        "respective 99th percentiles for display only; full extreme values are retained "
        "in the source data and reported in the descriptive tables."
    )

    c5 = find_caption(doc, "Fig. 5 High-frequency return behaviour")
    c5.text = (
        "Fig. 5 High-frequency return behaviour in the Bronze one-minute dataset. "
        "(a) One-minute log-return distribution; (b) log-return boxplot; (c) monthly "
        "log-return volatility; and (d) 60-minute rolling volatility. Panel (a) is "
        "displayed within ±P99.9 of the absolute one-minute log return to make the "
        "central distribution legible; tail observations are retained and reported "
        "separately in the extreme-event analysis."
    )

    doc.save(args.output_docx)
    print(args.output_docx)


if __name__ == "__main__":
    main()
